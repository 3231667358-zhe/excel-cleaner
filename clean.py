#!/usr/bin/env python3
"""
Excel / CSV 自动整理工具 V1

V1 只做三件事：
  1. 去首尾空白 —— 每个单元格（包括第一行的列名）两端的空格、换行都去掉
  2. 删除整行空白 —— 去空白后整行都是 None 或 "" 的行，直接删掉
  3. 整行去重   —— 内容完全相同的行，只保留最前面那一条

空单元格本身不会被填任何东西（不会填 0，也不会填 NULL）。

支持的文件：.csv / .txt / .xlsx
不支持 .xlsm —— 本工具会重建工作簿，宏无法安全保留。

Excel 额外有一道闸门：只接受「无公式的纯数据表」。
  发现真正的公式单元格 → 拒绝处理，不产出任何文件（删行不会平移公式引用，会算错）
  发现内容是 "=" 开头的文本单元格 → 拒绝处理（重写时会被误当成公式）
  CSV / TXT 不受这道闸门影响，行为完全不变。

安全约定（写死在代码里，不靠自觉）：
  - 永远不覆盖输入文件，结果写到另一个文件
  - 默认输出文件名是「原文件名_cleaned.原后缀」
  - 输出文件已经存在时，必须显式加 --force 才会覆盖

用法示例：
  python3 clean.py 原始数据.csv
  python3 clean.py 原始数据.xlsx -o 结果.xlsx
  python3 clean.py 原始数据.csv --no-dedup --report 报告.txt

看详细帮助：python3 clean.py -h
"""

import argparse
import csv
import sys
from pathlib import Path

EXCEL_SUFFIXES = {".xlsx"}
CSV_SUFFIXES = {".csv", ".txt"}

XLSM_HINT = (
    "不支持 .xlsm：本工具要重建工作簿，宏（VBA）没办法安全保留，硬做会静默丢宏。\n"
    "请先在 Excel 里「另存为 .xlsx」，再处理另存出来的文件。"
)

# 读 CSV 时按这个顺序试编码：先 UTF-8（含 BOM），再 GB18030（GBK 的超集，兼容老中文表格）
CSV_INPUT_ENCODINGS = ("utf-8-sig", "gb18030")
# 写 CSV 用带 BOM 的 UTF-8，这样 Windows 版 Excel 双击打开不会中文乱码
CSV_OUTPUT_ENCODING = "utf-8-sig"


class CleanError(Exception):
    """本工具能对用户说清楚的问题。捕获后只打印提示，不打印 Python 堆栈。"""


# --------------------------------------------------------------------------
# 核心清理逻辑
# --------------------------------------------------------------------------

def trim_cell(value):
    """去掉字符串两端的空白；数字、日期、空值原样返回。"""
    if isinstance(value, str):
        return value.strip()
    return value


def count_changed(before, after):
    """数一数这一行里有多少个单元格的内容变了。"""
    return sum(1 for old, new in zip(before, after) if old != new)


def is_blank_row(row):
    """
    判断一整行是不是空白行：每个单元格都是 None 或空字符串 ""。

    注意 0、0.0、"0" 都不算空，不会被删。空单元格也不会被填成任何东西。
    """
    return all(cell is None or cell == "" for cell in row)


def clean_table(rows, do_trim=True, do_dedup=True):
    """
    清理一张二维表。

    rows 是一行一行的列表，第一行被当作表头：
      - 表头永远只去空白，不参与去重，也不会被「删除空白行」规则删掉（即使它是空的）
    每行数据的处理顺序：去空白 → 是空白行就删 → 再判重
    返回 (清理后的行列表, 统计字典)

    注意：统计的空白只算真正留在输出里的行。被删掉的行（重复行、空白行）里面的空白
    不计数，免得报告里的数字和输出文件对不上。
    """
    stats = {
        "read_rows": len(rows),
        "trimmed": 0,
        "dropped": 0,
        "dropped_blank": 0,
        "kept": 0,
    }
    if not rows:
        return [], stats

    def fix_row(row):
        return [trim_cell(cell) if do_trim else cell for cell in row]

    header = fix_row(rows[0])
    cleaned = [header]
    if do_trim:
        stats["trimmed"] += count_changed(rows[0], header)

    # 表头不进 seen：它只去空白、不参与去重。
    # 否则「内容和表头一模一样的数据行」会被当成重复行删掉，而它其实是数据。
    seen = set()

    for row in rows[1:]:
        new_row = fix_row(row)

        # 规则一：整行空白就删掉
        if is_blank_row(new_row):
            stats["dropped_blank"] += 1
            continue

        # 规则二：整行去重
        if do_dedup:
            key = tuple(new_row)
            if key in seen:
                stats["dropped"] += 1
                continue
            seen.add(key)

        if do_trim:
            stats["trimmed"] += count_changed(row, new_row)
        cleaned.append(new_row)

    stats["kept"] = len(cleaned)
    return cleaned, stats


# --------------------------------------------------------------------------
# 读文件
# --------------------------------------------------------------------------

def read_csv_rows(path):
    """读 CSV，返回二维列表。会自动试 UTF-8 和 GB18030 两种编码。"""
    last_error = None
    for encoding in CSV_INPUT_ENCODINGS:
        try:
            with path.open("r", encoding=encoding, newline="") as f:
                return [row for row in csv.reader(f)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise CleanError(
        "读不了这个 CSV：按 {} 依次尝试都解码失败。\n原始报错：{}".format(
            " / ".join(CSV_INPUT_ENCODINGS), last_error
        )
    )


def load_openpyxl():
    """按需导入 openpyxl，没装的话给一条能直接照抄的命令。"""
    try:
        import openpyxl
    except ImportError:
        raise CleanError(
            "这是 Excel 文件，读写需要 openpyxl，但当前 Python 里没有这个库。\n"
            "  Ubuntu 安装（先跑这一条）：sudo apt install python3-openpyxl\n"
            "  装完重新运行本工具即可。只想先试 CSV 的话可以跳过。"
        )
    return openpyxl


def read_excel_workbook(path):
    """
    读一遍 Excel，同时干两件事：

      1. 把所有工作表转成二维列表
      2. 记下两类「危险单元格」：真正的公式单元格、内容是 "=" 开头的文本单元格

    返回 ([(表名, 二维列表), ...], 公式单元格列表, 危险文本单元格列表)

    这里故意不用 values_only，而是逐个 cell 对象读：只有 cell.data_type 才能区分
    「真正的公式」和「内容是 = 开头的纯文本」。两者的 cell.value 长得一模一样，
    但前者重写后还是公式，后者会被 openpyxl 误当成公式。
    """
    openpyxl = load_openpyxl()
    try:
        workbook = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        # openpyxl 的异常类型有好几种（文件损坏、其实是 .xls、被占用等），统一转成一句人话
        raise CleanError("打不开这个 Excel：{}".format(exc))
    try:
        tables = []
        formulas = []
        equal_texts = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    values.append(cell.value)
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        if cell.data_type == "f":
                            formulas.append((sheet.title, cell.coordinate, value))
                        else:
                            equal_texts.append((sheet.title, cell.coordinate, value))
                rows.append(values)
            tables.append((sheet.title, rows))
        return tables, formulas, equal_texts
    finally:
        workbook.close()


def format_hits(hits, limit=5):
    """把「危险单元格」列表排版成几行，最多列 5 个，避免刷屏。"""
    lines = [
        "    [{0}] {1}: {2}".format(sheet, coord, value)
        for sheet, coord, value in hits[:limit]
    ]
    if len(hits) > limit:
        lines.append("    …… 另外还有 {} 处".format(len(hits) - limit))
    return lines


def check_excel_is_plain_data(formulas, equal_texts):
    """
    Excel 安全闸门：只放行「无公式的纯数据表」。

    V1 不修公式，也不记录单元格原始类型，所以一旦发现这两类单元格就直接拒绝，
    宁可不做，也不给你一份算错的表。这个函数在写盘之前调用，拒绝时不会有任何输出文件。
    """
    if not formulas and not equal_texts:
        return

    parts = ["这个 Excel 不能处理，已取消，没有生成任何输出文件。"]

    if formulas:
        parts.append(
            "发现真正的公式单元格，共 {} 处：\n{}".format(
                len(formulas), "\n".join(format_hits(formulas))
            )
        )
        parts.append("V1 只支持无公式的纯数据表，为避免公式引用被破坏，本次处理已取消。")

    if equal_texts:
        parts.append(
            "发现原本是文本、但内容以 = 开头的单元格，共 {} 处：\n{}".format(
                len(equal_texts), "\n".join(format_hits(equal_texts))
            )
        )
        parts.append("为避免这些文本被误写成公式，本次处理已取消。")

    parts.append(
        "怎么改：在 Excel 里把公式选择性粘贴成数值（或先把那些 = 开头的文本改掉），"
        "另存为新的 .xlsx，再拿新文件来整理。"
    )
    raise CleanError("\n".join(parts))


# --------------------------------------------------------------------------
# 写文件
# --------------------------------------------------------------------------

def write_csv_rows(path, rows):
    with path.open("w", encoding=CSV_OUTPUT_ENCODING, newline="") as f:
        csv.writer(f).writerows(rows)


def write_excel_sheets(path, tables):
    """
    把 [(表名, 二维列表), ...] 写成一个新的 xlsx。

    这里是「重建工作簿」，所以只搬单元格内容：颜色、字体、列宽、合并单元格都不保留，
    宏更是无法保留 —— 这就是本工具不接 .xlsm 的原因。
    """
    openpyxl = load_openpyxl()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)  # 删掉默认那张空 Sheet，免得输出多一张空表
    for title, rows in tables:
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)


# --------------------------------------------------------------------------
# 参数、路径检查、主流程
# --------------------------------------------------------------------------

def build_parser():
    parser = argparse.ArgumentParser(
        description="Excel / CSV 自动整理工具 V1：去首尾空白 + 删除整行空白 + 整行去重，永不覆盖原始文件。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "示例：\n"
            "  python3 clean.py 原始数据.csv\n"
            "  python3 clean.py 原始数据.xlsx -o 结果.xlsx\n"
            "  python3 clean.py 原始数据.csv --no-dedup --report 报告.txt\n"
        ),
    )
    parser.add_argument("input", help="要整理的源文件（.csv / .txt / .xlsx）")
    parser.add_argument("-o", "--output", help="结果写到哪；默认：原文件名_cleaned.原后缀")
    parser.add_argument("--report", help="把本次统计结果额外写一份到这个文本文件")
    parser.add_argument("--no-trim", action="store_true", help="跳过「去首尾空白」")
    parser.add_argument("--no-dedup", action="store_true", help="跳过「整行去重」")
    parser.add_argument("--force", action="store_true", help="输出文件已存在时，允许覆盖它")
    return parser


def default_output_path(input_path):
    return input_path.with_name(input_path.stem + "_cleaned" + input_path.suffix)


def check_paths(input_path, output_path, report_path):
    """所有会写到磁盘之前的检查都放这里，任何一个不满足就直接报错退出。"""
    if not input_path.exists():
        raise CleanError("找不到输入文件：{}".format(input_path))
    if not input_path.is_file():
        raise CleanError("输入不是一个普通文件：{}".format(input_path))

    if output_path.resolve() == input_path.resolve():
        raise CleanError(
            "输出路径和输入文件是同一个文件。本工具不允许覆盖原始文件，请用 -o 换一个输出名。"
        )
    if report_path is not None and report_path.resolve() in (
        input_path.resolve(),
        output_path.resolve(),
    ):
        raise CleanError("--report 的路径和输入/输出文件重了，换一个路径。")


def detect_kind(path):
    """根据扩展名判断这是 CSV 还是 Excel，不认识就报错。"""
    suffix = path.suffix.lower()
    if suffix == ".xlsm":
        raise CleanError(XLSM_HINT)
    if suffix in EXCEL_SUFFIXES:
        return "excel", suffix
    if suffix in CSV_SUFFIXES:
        return "csv", suffix
    raise CleanError(
        "不认识的扩展名：{}\nV1 支持的输入：{}".format(
            suffix or "（没有扩展名）",
            "、".join(sorted(EXCEL_SUFFIXES | CSV_SUFFIXES)),
        )
    )


def format_stats(stats):
    """把一张表的统计数字拼成一句人话，终端和 --report 共用。"""
    return (
        "读入 {read_rows} 行 → 输出 {kept} 行；去掉空白 {trimmed} 处，"
        "删除重复 {dropped} 行，删除空白行 {dropped_blank} 行"
    ).format(**stats)


def build_summary(input_path, output_path, kind, do_trim, do_dedup, total, detail_lines):
    """拼出给人看的统计文本（终端打印和 --report 用的是同一份）。"""
    rules = [
        "去首尾空白" + ("" if do_trim else "（已关闭）"),
        "删除整行空白",
        "整行去重" + ("" if do_dedup else "（已关闭）"),
    ]
    lines = [
        "整理完成",
        "  源文件：{}".format(input_path),
        "  结果文件：{}".format(output_path),
        "  启用规则：{}".format("、".join(rules)),
        "  汇总：{}".format(format_stats(total)),
    ]
    if kind == "excel":
        lines.append("  分表统计：")
        lines.extend(detail_lines)
    if total["read_rows"] == 0:
        lines.append("  注意：源文件里没有读到任何行，检查一下文件是不是空的。")
    return "\n".join(lines)


def run(input_path, output_path, report_path, do_trim, do_dedup):
    kind, _ = detect_kind(input_path)

    # 输出格式必须和输入一致，避免出现「xlsx 写进 .csv」这种看不懂的结果
    output_suffix = output_path.suffix.lower()
    allowed_suffixes = EXCEL_SUFFIXES if kind == "excel" else CSV_SUFFIXES
    if output_suffix not in allowed_suffixes:
        if output_suffix == ".xlsm":
            raise CleanError(XLSM_HINT)
        raise CleanError(
            "输入是 {}，输出也应该是 {} 里的后缀，现在给的是 {}。本工具不做格式转换。".format(
                input_path.suffix,
                "、".join(sorted(allowed_suffixes)),
                output_suffix or "（空）",
            )
        )

    # 先把文件读进来并做完安全检查，任何拒绝都必须发生在「建目录 / 写文件」之前，
    # 这样被拒绝时磁盘上不会多出任何东西。
    if kind == "csv":
        tables = [(input_path.stem, read_csv_rows(input_path))]
    else:
        tables, formulas, equal_texts = read_excel_workbook(input_path)
        check_excel_is_plain_data(formulas, equal_texts)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    cleaned_tables = []
    total = {"read_rows": 0, "trimmed": 0, "dropped": 0, "dropped_blank": 0, "kept": 0}
    detail_lines = []
    for title, rows in tables:
        cleaned, stats = clean_table(rows, do_trim=do_trim, do_dedup=do_dedup)
        cleaned_tables.append((title, cleaned))
        for key in total:
            total[key] += stats[key]
        detail_lines.append("    [{}] {}".format(title, format_stats(stats)))

    if kind == "csv":
        write_csv_rows(output_path, cleaned_tables[0][1])
    else:
        write_excel_sheets(output_path, cleaned_tables)

    summary = build_summary(
        input_path, output_path, kind, do_trim, do_dedup, total, detail_lines
    )
    print(summary)

    if report_path is not None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(summary + "\n", encoding="utf-8")
        print("\n报告已写入：{}".format(report_path))


def main(argv=None):
    args = build_parser().parse_args(argv)

    input_path = Path(args.input).expanduser()
    output_path = (
        Path(args.output).expanduser() if args.output else default_output_path(input_path)
    )
    report_path = Path(args.report).expanduser() if args.report else None

    try:
        check_paths(input_path, output_path, report_path)
        if output_path.exists() and not args.force:
            raise CleanError(
                "输出文件已存在：{}\n"
                "本工具默认不覆盖任何已有文件。确实要覆盖就加 --force，或者换个 -o 路径。".format(
                    output_path
                )
            )
        run(
            input_path,
            output_path,
            report_path,
            do_trim=not args.no_trim,
            do_dedup=not args.no_dedup,
        )
    except CleanError as exc:
        print("错误：{}".format(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
